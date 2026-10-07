import requests
import io
import re
import openpyxl

session = requests.Session()
BASE = 'http://127.0.0.1:5000'

# 1. Unauthenticated test
r = session.get(f'{BASE}/', allow_redirects=False)
assert r.status_code == 302 and '/login' in r.headers.get('Location', ''), f'Failed auth redirect: {r.status_code}'
print('[PASS] Unauthenticated redirects to /login')

# 2. Fresh signup
signup_payload = {
    'name': 'Dr. Elena Rostova',
    'username': 'elena_rostova',
    'password': 'PasswordElena2026!',
    'classes_input': 'Grade 10 - Div A, Grade 10 - Div B, Grade 9 - Div C'
}
r = session.post(f'{BASE}/signup', data=signup_payload, allow_redirects=True)
assert r.status_code == 200
assert 'No Records Uploaded for' in r.text
assert 'Grade 10 - Div A' in r.text
# Verify zero dummy data pre-loaded
assert 'Alice' not in r.text and 'Bob' not in r.text
print('[PASS] Fresh signup redirects to clean dashboard with ZERO dummy data')

# 3. Create Multi-Division CSV test files (Suggestion 1.2)
att_csv = '''StudentID,Name,Division,Date,Status
201,Sophia Clark,A,2026-09-01,Present
202,Lucas Miller,A,2026-09-01,Present
203,Oliver Garcia,B,2026-09-01,Absent
204,Emma Davis,B,2026-09-01,Present
201,Sophia Clark,A,2026-09-02,Present
202,Lucas Miller,A,2026-09-02,Absent
203,Oliver Garcia,B,2026-09-02,Present
204,Emma Davis,B,2026-09-02,Present
'''

marks_csv = '''StudentID,Name,Division,Subject,Marks,MaxMarks
201,Sophia Clark,A,Mathematics,95,100
201,Sophia Clark,A,Science,92,100
202,Lucas Miller,A,Mathematics,60,100
202,Lucas Miller,A,Science,65,100
203,Oliver Garcia,B,Mathematics,35,100
203,Oliver Garcia,B,Science,40,100
204,Emma Davis,B,Mathematics,88,100
204,Emma Davis,B,Science,85,100
'''

files = {
    'attendance_file': ('attendance_multidiv.csv', io.BytesIO(att_csv.encode('utf-8')), 'text/csv'),
    'marks_file': ('marks_multidiv.csv', io.BytesIO(marks_csv.encode('utf-8')), 'text/csv')
}
data = {
    'selected_class': 'Grade 10 - Div A',
    'attendance_threshold': '75.0',
    'pass_percentage': '40.0'
}

r = session.post(f'{BASE}/generate', files=files, data=data, allow_redirects=True)
assert r.status_code == 200
print('[PASS] Report generated for Grade 10 - Div A')

# Check filename pattern
assert 'Grade_10_Div_A_Report_' in r.text
print('[PASS] Clean organized Excel filename generated')

# Verify side-by-side division comparison on dashboard
assert 'Division Comparison' in r.text
assert 'Div A' in r.text and 'Div B' in r.text
assert 'Sophia Clark' in r.text
print('[PASS] Dashboard displays division filter pills and side-by-side comparison')

# Download the generated excel file and inspect sheet names
match = re.search(r'href="/download/([^"]+)"', r.text)
assert match, 'Download link not found'
filename = match.group(1)
print(f'[INFO] Downloading {filename}')

r_dl = session.get(f'{BASE}/download/{filename}')
wb = openpyxl.load_workbook(io.BytesIO(r_dl.content))
print(f'[INFO] Excel sheet names: {wb.sheetnames}')
assert 'Master_Summary' in wb.sheetnames
assert 'Attendance_DivA' in wb.sheetnames
assert 'Attendance_DivB' in wb.sheetnames
assert 'Marks_DivA' in wb.sheetnames
assert 'Marks_DivB' in wb.sheetnames
print('[PASS] Excel workbook contains Master_Summary and dedicated Division sub-sheets!')

# 4. Switch to Grade 10 - Div B
r_switch = session.post(f'{BASE}/switch-class', data={'active_class': 'Grade 10 - Div B'}, allow_redirects=True)
r_b = session.get(f'{BASE}/dashboard')
assert 'No Records Uploaded for' in r_b.text
assert 'Grade 10 - Div B' in r_b.text
print('[PASS] Switching to Grade 10 - Div B shows clean empty state')

# 5. Switch back to Grade 10 - Div A
r_switch_back = session.post(f'{BASE}/switch-class', data={'active_class': 'Grade 10 - Div A'}, allow_redirects=True)
r_a = session.get(f'{BASE}/dashboard')
assert 'Sophia Clark' in r_a.text
print('[PASS] Switching back to Grade 10 - Div A restores its records')

# 6. Test templates and settings pages
r_tmpl = session.get(f'{BASE}/templates')
assert r_tmpl.status_code == 200
r_sett = session.get(f'{BASE}/settings')
assert r_sett.status_code == 200
print('[PASS] /templates and /settings load successfully')

# 7. Test Logout and Login
session.get(f'{BASE}/logout')
r_unauth = session.get(f'{BASE}/dashboard', allow_redirects=False)
assert r_unauth.status_code == 302
print('[PASS] Logout successfully clears session')

r_login = session.post(f'{BASE}/login', data={'username': 'elena_rostova', 'password': 'PasswordElena2026!'}, allow_redirects=True)
assert r_login.status_code == 200
assert 'Academic Performance' in r_login.text
print('[PASS] Login succeeds and restores teacher dashboard')

print('\nALL TESTS PASSED WITH 100% SUCCESS!')
