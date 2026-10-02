import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import PatternFill
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule

# ---------------------------------------------------------
# CONSTANTS & STYLES
# ---------------------------------------------------------
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

def _get_column_letter(ws, header_row_idx: int, column_name: str) -> str:
    """
    Find the Excel column letter for a given DataFrame column name,
    by scanning the header row of the worksheet.
    """
    for cell in ws[header_row_idx]:
        if cell.value == column_name:
            return cell.column_letter
    return None

def _add_attendance_chart(ws_summary, ws_attendance, attendance_df):
    """
    Build a BarChart for Attendance sourcing data from the Attendance sheet.
    """
    if attendance_df.empty or ws_attendance.max_row <= 1:
        return
        
    name_col_idx = None
    att_col_idx = None
    
    # Find the column indices for "Name" and "AttendancePercent"
    for cell in ws_attendance[1]:
        if cell.value == "Name":
            name_col_idx = cell.column
        elif cell.value == "AttendancePercent":
            att_col_idx = cell.column
            
    if not name_col_idx or not att_col_idx:
        return

    max_row = ws_attendance.max_row
    
    chart = BarChart()
    chart.title = "Attendance by Student"
    chart.width = 12
    chart.height = 7
    
    # Values: AttendancePercent (including header)
    data = Reference(ws_attendance, min_col=att_col_idx, min_row=1, max_row=max_row)
    # Categories: Name (excluding header)
    cats = Reference(ws_attendance, min_col=name_col_idx, min_row=2, max_row=max_row)
    
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    ws_summary.add_chart(chart, "A1")


def _add_grade_chart(ws_summary, marks_df):
    """
    Build a BarChart for Grade Distribution by writing a small table to the Summary sheet.
    """
    if marks_df.empty:
        return
        
    grade_counts = marks_df['Grade'].value_counts().sort_index()
    
    start_row = 20
    ws_summary.cell(row=start_row, column=1, value="Grade")
    ws_summary.cell(row=start_row, column=2, value="Count")
    
    current_row = start_row + 1
    for grade, count in grade_counts.items():
        ws_summary.cell(row=current_row, column=1, value=grade)
        ws_summary.cell(row=current_row, column=2, value=count)
        current_row += 1
        
    max_row = current_row - 1
    
    chart = BarChart()
    chart.title = "Grade Distribution"
    chart.width = 12
    chart.height = 7
    
    data = Reference(ws_summary, min_col=2, min_row=start_row, max_row=max_row)
    cats = Reference(ws_summary, min_col=1, min_row=start_row+1, max_row=max_row)
    
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    # Position so it doesn't overlap attendance chart
    ws_summary.add_chart(chart, "I1")


def _add_pass_fail_chart(ws_summary, marks_df):
    """
    Build a PieChart for Pass vs Fail by writing a small table to the Summary sheet.
    """
    if marks_df.empty:
        return
        
    fails = (marks_df['Grade'] == 'F').sum()
    passes = len(marks_df) - fails
    
    start_row = 30
    ws_summary.cell(row=start_row, column=1, value="Status")
    ws_summary.cell(row=start_row, column=2, value="Count")
    ws_summary.cell(row=start_row+1, column=1, value="Pass")
    ws_summary.cell(row=start_row+1, column=2, value=passes)
    ws_summary.cell(row=start_row+2, column=1, value="Fail")
    ws_summary.cell(row=start_row+2, column=2, value=fails)
    
    chart = PieChart()
    chart.title = "Pass vs Fail"
    chart.width = 12
    chart.height = 7
    
    data = Reference(ws_summary, min_col=2, min_row=start_row, max_row=start_row+2)
    cats = Reference(ws_summary, min_col=1, min_row=start_row+1, max_row=start_row+2)
    
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    
    # Position so it doesn't overlap other charts
    ws_summary.add_chart(chart, "A16")


def generate_report(attendance_df: pd.DataFrame, marks_df: pd.DataFrame, output_path: str) -> None:
    """
    Generate a consolidated Excel report from attendance and marks DataFrames.
    Includes a Summary dashboard sheet, data sheets, and conditional formatting.
    """
    wb = Workbook()
    
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    # 1. Create a "Summary" sheet and insert it at index=0
    ws_summary = wb.create_sheet("Summary", 0)

    # ---------------------------------------------------------
    # ATTENDANCE SHEET
    # ---------------------------------------------------------
    ws_attendance = wb.create_sheet(title="Attendance")
    for row in dataframe_to_rows(attendance_df, index=False, header=True):
        ws_attendance.append(row)
        
    if ws_attendance.max_row > 1:
        att_col = _get_column_letter(ws_attendance, 1, "AttendancePercent")
        if att_col:
            rule = CellIsRule(operator='lessThan', formula=['75'], stopIfTrue=True, fill=RED_FILL)
            ws_attendance.conditional_formatting.add(f'{att_col}2:{att_col}{ws_attendance.max_row}', rule)

    # ---------------------------------------------------------
    # MARKS SHEET
    # ---------------------------------------------------------
    ws_marks = wb.create_sheet(title="Marks")
    for row in dataframe_to_rows(marks_df, index=False, header=True):
        ws_marks.append(row)
        
    if ws_marks.max_row > 1:
        grade_col = _get_column_letter(ws_marks, 1, "Grade")
        if grade_col:
            rule = CellIsRule(operator='equal', formula=['"F"'], stopIfTrue=True, fill=RED_FILL)
            ws_marks.conditional_formatting.add(f'{grade_col}2:{grade_col}{ws_marks.max_row}', rule)

    # ---------------------------------------------------------
    # CHARTS & DASHBOARDS (Week 3)
    # ---------------------------------------------------------
    _add_attendance_chart(ws_summary, ws_attendance, attendance_df)
    _add_grade_chart(ws_summary, marks_df)
    _add_pass_fail_chart(ws_summary, marks_df)

    # 5. Save the workbook to output_path
    wb.save(output_path)
    print(f"Report successfully saved to {output_path}")


if __name__ == "__main__":
    # ---------------------------------------------------------
    # DUMMY DATA FOR TESTING
    # ---------------------------------------------------------
    dummy_attendance_data = {
        'StudentID': [101, 102, 103, 104, 105, 106],
        'Name': ['Alice', 'Bob', 'Charlie', 'Diana', 'Eve', 'Frank'],
        'AttendancePercent': [85.5, 60.0, 92.0, 72.5, 100.0, 45.0]  # Bob, Diana, Frank < 75
    }
    df_attendance = pd.DataFrame(dummy_attendance_data)
    
    dummy_marks_data = {
        'StudentID': [101, 102, 103, 104, 105, 106],
        'Name': ['Alice', 'Bob', 'Charlie', 'Diana', 'Eve', 'Frank'],
        'Subject': ['Math', 'Math', 'Math', 'Math', 'Math', 'Math'],
        'Marks': [95, 35, 75, 82, 88, 40],
        'MaxMarks': [100, 100, 100, 100, 100, 100],
        'Grade': ['A+', 'F', 'B', 'A', 'A+', 'F']  # Bob, Frank have F
    }
    df_marks = pd.DataFrame(dummy_marks_data)

    test_output_file = "test_output.xlsx"
    print("Generating Week 3 test report with Summary charts...")
    generate_report(df_attendance, df_marks, test_output_file)
    print("Done! Check your folder for test_output.xlsx to verify the Summary tab and charts.")
